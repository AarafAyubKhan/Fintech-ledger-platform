"""
FinSight AI — DOCX Parser
Extracts text from Word documents.
"""

from docx import Document
import structlog

logger = structlog.get_logger()


class DocxParser:
    """Parse DOCX documents."""

    async def parse(self, file_path: str) -> tuple[str, int]:
        """Extract text from a DOCX file."""
        try:
            doc = Document(file_path)
            paragraphs = []
            page_estimate = 1

            for para in doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text)

            # Extract table content
            for table in doc.tables:
                table_rows = []
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    table_rows.append(" | ".join(cells))
                if table_rows:
                    paragraphs.append("[Table]\n" + "\n".join(table_rows))

            full_text = "\n\n".join(paragraphs)
            # Rough page estimation (250 words per page)
            page_estimate = max(1, len(full_text.split()) // 250)

            logger.info("docx_parsed", file=file_path, chars=len(full_text))
            return full_text, page_estimate

        except Exception as e:
            logger.error("docx_parse_error", file=file_path, error=str(e))
            raise
